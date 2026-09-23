import openpyxl
import sys

def main():
    sys.stdout.reconfigure(encoding='utf-8')
    wb = openpyxl.load_workbook('02_Data/Optimierung des Structural Equivalency Spreadsheets_Schnee_2.xlsm', data_only=True, read_only=True)
    ws = wb['Overview']
    for r in range(5, 60):
        val = ws.cell(r, 1).value
        if val in ['27_01', '27_02', '27_03', '27_04', '26_16', '26_17', '26_18', '26_19']:
            row_vals = [ws.cell(r, c).value for c in range(1, 30)]
            print(f"Row {r:2d} ({val}): {row_vals}")
            
    print("\n--- Inspecting Sheet 27_04 ---")
    ws4 = wb['27_04']
    for r in range(1, 32):
        row_vals = [ws4.cell(r, c).value for c in range(1, 12)]
        if any(v is not None for v in row_vals):
            print(f"27_04 R{r:2d}: {row_vals}")
            
    print("\n--- Inspecting Sheet 27_03 ---")
    ws3 = wb['27_03']
    for r in range(1, 32):
        row_vals = [ws3.cell(r, c).value for c in range(1, 12)]
        if any(v is not None for v in row_vals):
            print(f"27_03 R{r:2d}: {row_vals}")

if __name__ == '__main__':
    main()
