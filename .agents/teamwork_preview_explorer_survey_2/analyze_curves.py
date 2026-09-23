import openpyxl
import sys

def to_float(val):
    if val is None:
        return None
    if isinstance(val, (int, float)):
        return float(val)
    try:
        s = str(val).strip().replace(',', '.')
        return float(s)
    except Exception:
        return None

def analyze_sheet_curve(ws):
    max_d = 0.0
    f_at_max_d = 0.0
    max_f = 0.0
    d_at_max_f = 0.0
    f_at_12_7 = None
    
    energy_12_7 = 0.0
    total_energy = 0.0
    
    pts = []
    # iter_rows is streaming in read_only mode
    for row in ws.iter_rows(min_row=32, values_only=True):
        if not row or len(row) < 4:
            continue
        c_d = to_float(row[1]) # col B (0-based idx 1)
        if c_d is None and len(row) > 2:
            c_d = to_float(row[2]) # col C
        c_f = to_float(row[3]) # col D (Kraft N)
        if c_d is None or c_f is None:
            continue
        d = abs(c_d)
        f = abs(c_f)
        pts.append((d, f))
    
    if not pts:
        return None
        
    for i in range(len(pts)):
        d, f = pts[i]
        if f > max_f:
            max_f = f
            d_at_max_f = d
        if d > max_d:
            max_d = d
            f_at_max_d = f
            
        if i > 0:
            d_prev, f_prev = pts[i-1]
            delta_d = d - d_prev
            if delta_d > 0:
                trapz = 0.5 * (f + f_prev) * delta_d / 1000.0 # N*mm -> J
                total_energy += trapz
                if d <= 12.7:
                    energy_12_7 += trapz
                elif d_prev < 12.7:
                    frac = (12.7 - d_prev) / delta_d
                    f_12_7 = f_prev + frac * (f - f_prev)
                    energy_12_7 += 0.5 * (f_prev + f_12_7) * (12.7 - d_prev) / 1000.0
                    f_at_12_7 = f_12_7
                    
    if f_at_12_7 is None and max_d >= 12.7:
        for i in range(len(pts)-1):
            if pts[i][0] <= 12.7 <= pts[i+1][0]:
                d0, f0 = pts[i]
                d1, f1 = pts[i+1]
                if d1 > d0:
                    f_at_12_7 = f0 + (12.7 - d0)/(d1 - d0) * (f1 - f0)
                break
                
    return {
        'num_points': len(pts),
        'max_f': max_f,
        'd_at_max_f': d_at_max_f,
        'max_d': max_d,
        'f_at_max_d': f_at_max_d,
        'f_at_12_7': f_at_12_7,
        'energy_12_7': energy_12_7 if max_d >= 12.7 else None,
        'total_energy': total_energy
    }

def main():
    sys.stdout.reconfigure(encoding='utf-8')
    wb = openpyxl.load_workbook('02_Data/Optimierung des Structural Equivalency Spreadsheets_Schnee_2.xlsm', data_only=True, read_only=True)
    sheets_to_check = ['26_17', '26_18', '26_19', '26_11', '26_12', '27_01', '27_02', '27_03', '27_04',
                       '24_02_3P', '24_03_3P', '24_05_3P', '24_07_3P', '24_11_3P', '24_12_3P', '24_13_3P']
    
    print(f"{'Sheet':<12} | {'Fmax [N]':<10} | {'d@Fmax [mm]':<11} | {'max_d [mm]':<10} | {'F@max_d [N]':<11} | {'F@12.7mm':<10} | {'E@12.7 [J]':<11} | {'E_tot [J]':<10}")
    print("-" * 100)
    for s in sheets_to_check:
        if s not in wb.sheetnames:
            print(f"{s:<12} | NOT FOUND")
            continue
        ws = wb[s]
        res = analyze_sheet_curve(ws)
        if res:
            e127_str = f"{res['energy_12_7']:.1f}" if res['energy_12_7'] is not None else "N/A (<12.7)"
            f127_str = f"{res['f_at_12_7']:.1f}" if res['f_at_12_7'] is not None else "N/A"
            print(f"{s:<12} | {res['max_f']:<10.1f} | {res['d_at_max_f']:<11.2f} | {res['max_d']:<10.2f} | {res['f_at_max_d']:<11.1f} | {f127_str:<10} | {e127_str:<11} | {res['total_energy']:<10.1f}")
        else:
            print(f"{s:<12} | NO CURVE DATA")

if __name__ == '__main__':
    main()
