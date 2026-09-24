import openpyxl
import sys

def main():
    sys.stdout.reconfigure(encoding='utf-8')
    wb = openpyxl.load_workbook('02_Data/Optimierung des Structural Equivalency Spreadsheets_Schnee_2.xlsm', data_only=True, read_only=True)
    
    for s in ['Auslegung neuer Panels', 'Materialien', 'Akku_MaterialData', 'MaterialData']:
        if s not in wb.sheetnames:
            continue
        ws = wb[s]
        print(f"\n{'='*50}\nSheet: {s}\n{'='*50}")
        count = 0
        for r_idx, row in enumerate(ws.iter_rows(values_only=True), start=1):
            if any(v is not None for v in row):
                print(f"R{r_idx:2d}: {[v for v in row[:12] if v is not None]}")
                count += 1
                if count >= 30:
                    break

if __name__ == '__main__':
    main()
